from datetime import date
from decimal import Decimal
from io import BytesIO

from flask import (
    Blueprint,
    render_template,
    request,
    send_file,
    abort
)
from flask_login import login_required, current_user
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.models import (
    ProduccionCafe,
    ProduccionPlatano,
    Venta,
    Inventario,
    MovimientoInventario
)


reportes_bp = Blueprint(
    "reportes",
    __name__,
    url_prefix="/reportes"
)


def obtener_finca_activa():
    if not current_user.is_authenticated or not current_user.finca:
        abort(403)

    return current_user.finca


def decimal_seguro(valor):
    if valor is None:
        return Decimal("0")

    return Decimal(str(valor))


def suma_decimal(valores):
    total = Decimal("0")

    for valor in valores:
        total += decimal_seguro(valor)

    return total


def obtener_periodo():
    hoy = date.today()

    anio_actual = hoy.year

    try:
        anio = int(request.args.get("anio", anio_actual))
    except ValueError:
        anio = anio_actual

    mes_raw = request.args.get("mes", "").strip()

    mes = None

    if mes_raw:
        try:
            mes = int(mes_raw)
        except ValueError:
            mes = None

    if mes and 1 <= mes <= 12:
        fecha_inicio = date(anio, mes, 1)

        if mes == 12:
            fecha_fin = date(anio + 1, 1, 1)
        else:
            fecha_fin = date(anio, mes + 1, 1)

        periodo_label = f"{nombre_mes(mes)} de {anio}"

    else:
        mes = None
        fecha_inicio = date(anio, 1, 1)
        fecha_fin = date(anio + 1, 1, 1)
        periodo_label = f"Año {anio}"

    return anio, mes, fecha_inicio, fecha_fin, periodo_label


def nombre_mes(mes):
    meses = {
        1: "Enero",
        2: "Febrero",
        3: "Marzo",
        4: "Abril",
        5: "Mayo",
        6: "Junio",
        7: "Julio",
        8: "Agosto",
        9: "Septiembre",
        10: "Octubre",
        11: "Noviembre",
        12: "Diciembre"
    }

    return meses.get(mes, "")


def texto_capitalizado(valor):
    if not valor:
        return ""

    return str(valor).replace("_", " ").strip().capitalize()


def valor_movimiento(movimiento):
    cantidad = decimal_seguro(movimiento.cantidad)

    costo_unitario = movimiento.costo_unitario

    if costo_unitario is None and movimiento.item:
        costo_unitario = movimiento.item.costo_unitario

    costo_unitario = decimal_seguro(costo_unitario)

    return cantidad * costo_unitario


def obtener_datos_reporte(finca_id, fecha_inicio, fecha_fin):
    producciones_cafe = (
        ProduccionCafe.query
        .filter_by(finca_id=finca_id)
        .filter(
            ProduccionCafe.fecha >= fecha_inicio,
            ProduccionCafe.fecha < fecha_fin
        )
        .order_by(ProduccionCafe.fecha.desc(), ProduccionCafe.id.desc())
        .all()
    )

    producciones_platano = (
        ProduccionPlatano.query
        .filter_by(finca_id=finca_id)
        .filter(
            ProduccionPlatano.fecha >= fecha_inicio,
            ProduccionPlatano.fecha < fecha_fin
        )
        .order_by(ProduccionPlatano.fecha.desc(), ProduccionPlatano.id.desc())
        .all()
    )

    ventas = (
        Venta.query
        .filter_by(finca_id=finca_id)
        .filter(
            Venta.fecha >= fecha_inicio,
            Venta.fecha < fecha_fin
        )
        .order_by(Venta.fecha.desc(), Venta.id.desc())
        .all()
    )

    movimientos = (
        MovimientoInventario.query
        .filter_by(finca_id=finca_id)
        .filter(
            MovimientoInventario.fecha >= fecha_inicio,
            MovimientoInventario.fecha < fecha_fin
        )
        .order_by(
            MovimientoInventario.fecha.desc(),
            MovimientoInventario.id.desc()
        )
        .all()
    )

    existencias = (
        Inventario.query
        .filter_by(finca_id=finca_id)
        .order_by(Inventario.categoria.asc(), Inventario.nombre.asc())
        .all()
    )

    filas_produccion = []

    for produccion in producciones_cafe:
        kilos = decimal_seguro(produccion.kilos_totales)
        arrobas = kilos / Decimal("12") if kilos > 0 else Decimal("0")

        filas_produccion.append({
            "fecha": produccion.fecha,
            "tipo_produccion": "Café",
            "producto": texto_capitalizado(produccion.tipo_cafe),
            "lote": produccion.lote,
            "kilos": kilos,
            "arrobas": arrobas,
            "racimos": "",
            "kg_promedio_racimo": "",
            "observaciones": produccion.observaciones or ""
        })

    for produccion in producciones_platano:
        kilos = decimal_seguro(produccion.kilos)
        racimos = produccion.racimos or 0

        if racimos > 0:
            kg_promedio = kilos / Decimal(racimos)
        else:
            kg_promedio = Decimal("0")

        filas_produccion.append({
            "fecha": produccion.fecha,
            "tipo_produccion": "Plátano / Banano / Comino",
            "producto": texto_capitalizado(produccion.tipo_cultivo),
            "lote": produccion.lote,
            "kilos": kilos,
            "arrobas": "",
            "racimos": racimos,
            "kg_promedio_racimo": kg_promedio,
            "observaciones": produccion.observaciones or ""
        })

    filas_produccion.sort(
        key=lambda fila: fila["fecha"],
        reverse=True
    )

    filas_ventas = []

    for venta in ventas:
        kilos = decimal_seguro(venta.kilos)
        racimos = venta.racimos or ""

        if venta.producto == "cafe":
            producto = "Café"
            arrobas = venta.arrobas or Decimal("0")
            kg_promedio_racimo = ""
        else:
            producto = "Plátano / Banano / Comino"
            arrobas = ""

            if venta.racimos and venta.racimos > 0:
                kg_promedio_racimo = kilos / Decimal(venta.racimos)
            else:
                kg_promedio_racimo = ""

        valor_esperado = decimal_seguro(venta.valor_calculado)
        valor_recibido = decimal_seguro(venta.valor_total)
        diferencia = decimal_seguro(venta.diferencia_valor)

        filas_ventas.append({
            "fecha": venta.fecha,
            "producto": producto,
            "tipo": venta.tipo_producto,
            "kilos": kilos,
            "arrobas": arrobas,
            "racimos": racimos,
            "kg_promedio_racimo": kg_promedio_racimo,
            "precio_unitario": decimal_seguro(venta.precio_unitario),
            "valor_esperado": valor_esperado,
            "valor_recibido": valor_recibido,
            "diferencia": diferencia,
            "observaciones": venta.observaciones or ""
        })

    filas_movimientos = []

    for movimiento in movimientos:
        item = movimiento.item

        cantidad = decimal_seguro(movimiento.cantidad)
        costo_unitario = movimiento.costo_unitario

        if costo_unitario is None and item:
            costo_unitario = item.costo_unitario

        costo_unitario = decimal_seguro(costo_unitario)
        valor_total = cantidad * costo_unitario

        filas_movimientos.append({
            "fecha": movimiento.fecha,
            "elemento": item.nombre if item else "",
            "categoria": item.categoria if item else "",
            "tipo_movimiento": texto_capitalizado(movimiento.tipo_movimiento),
            "cantidad": cantidad,
            "unidad": item.unidad_medida if item else "",
            "costo_unitario": costo_unitario,
            "valor_movimiento": valor_total,
            "proveedor": movimiento.proveedor or (item.proveedor if item else "") or "",
            "motivo": movimiento.motivo or "",
            "observaciones": movimiento.observaciones or ""
        })

    filas_existencias = []

    for item in existencias:
        cantidad = decimal_seguro(item.cantidad_disponible)
        costo_unitario = decimal_seguro(item.costo_unitario)
        valor_estimado = cantidad * costo_unitario

        filas_existencias.append({
            "elemento": item.nombre,
            "categoria": item.categoria,
            "cantidad_disponible": cantidad,
            "unidad": item.unidad_medida,
            "costo_unitario": costo_unitario,
            "valor_estimado": valor_estimado,
            "proveedor": item.proveedor or "",
            "ubicacion": item.ubicacion or "",
            "observaciones": item.observaciones or ""
        })

    total_cafe_kg = suma_decimal(
        produccion.kilos_totales
        for produccion in producciones_cafe
    )

    total_cafe_arrobas = (
        total_cafe_kg / Decimal("12")
        if total_cafe_kg > 0
        else Decimal("0")
    )

    total_agricola_kg = suma_decimal(
        produccion.kilos
        for produccion in producciones_platano
    )

    total_racimos = sum(
        produccion.racimos or 0
        for produccion in producciones_platano
    )

    total_ventas = suma_decimal(
        venta.valor_total
        for venta in ventas
    )

    ventas_cafe = suma_decimal(
        venta.valor_total
        for venta in ventas
        if venta.producto == "cafe"
    )

    ventas_agricolas = total_ventas - ventas_cafe

    valor_esperado_ventas = suma_decimal(
        venta.valor_calculado
        for venta in ventas
    )

    diferencia_ventas = total_ventas - valor_esperado_ventas

    total_entradas_valor = suma_decimal(
        valor_movimiento(movimiento)
        for movimiento in movimientos
        if movimiento.tipo_movimiento == "entrada"
    )

    total_salidas_valor = suma_decimal(
        valor_movimiento(movimiento)
        for movimiento in movimientos
        if movimiento.tipo_movimiento == "salida"
    )

    total_ajustes = sum(
        1
        for movimiento in movimientos
        if movimiento.tipo_movimiento == "ajuste"
    )

    valor_estimado_inventario = suma_decimal(
        item.valor_estimado
        for item in existencias
    )

    resumen = {
        "total_cafe_kg": total_cafe_kg,
        "total_cafe_arrobas": total_cafe_arrobas,
        "total_agricola_kg": total_agricola_kg,
        "total_racimos": total_racimos,

        "total_ventas": total_ventas,
        "ventas_cafe": ventas_cafe,
        "ventas_agricolas": ventas_agricolas,
        "valor_esperado_ventas": valor_esperado_ventas,
        "diferencia_ventas": diferencia_ventas,

        "cantidad_movimientos": len(movimientos),
        "total_entradas_valor": total_entradas_valor,
        "total_salidas_valor": total_salidas_valor,
        "total_ajustes": total_ajustes,

        "total_elementos_inventario": len(existencias),
        "valor_estimado_inventario": valor_estimado_inventario
    }

    return {
        "resumen": resumen,
        "filas_produccion": filas_produccion,
        "filas_ventas": filas_ventas,
        "filas_movimientos": filas_movimientos,
        "filas_existencias": filas_existencias
    }


@reportes_bp.route("/")
@login_required
def index():
    finca_activa = obtener_finca_activa()

    anio, mes, fecha_inicio, fecha_fin, periodo_label = obtener_periodo()

    datos = obtener_datos_reporte(
        finca_activa.id,
        fecha_inicio,
        fecha_fin
    )

    anio_actual = date.today().year

    anios = list(range(anio_actual - 5, anio_actual + 2))

    meses = [
        {"valor": 1, "nombre": "Enero"},
        {"valor": 2, "nombre": "Febrero"},
        {"valor": 3, "nombre": "Marzo"},
        {"valor": 4, "nombre": "Abril"},
        {"valor": 5, "nombre": "Mayo"},
        {"valor": 6, "nombre": "Junio"},
        {"valor": 7, "nombre": "Julio"},
        {"valor": 8, "nombre": "Agosto"},
        {"valor": 9, "nombre": "Septiembre"},
        {"valor": 10, "nombre": "Octubre"},
        {"valor": 11, "nombre": "Noviembre"},
        {"valor": 12, "nombre": "Diciembre"}
    ]

    return render_template(
        "reportes/index.html",
        finca_activa=finca_activa,
        anio=anio,
        mes=mes,
        anios=anios,
        meses=meses,
        periodo_label=periodo_label,
        resumen=datos["resumen"],
        filas_produccion=datos["filas_produccion"],
        filas_ventas=datos["filas_ventas"],
        filas_movimientos=datos["filas_movimientos"],
        filas_existencias=datos["filas_existencias"]
    )


def configurar_hoja(hoja):
    hoja.freeze_panes = "A2"

    for row in hoja.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

    for columna in hoja.columns:
        max_length = 0
        letra_columna = get_column_letter(columna[0].column)

        for cell in columna:
            value = cell.value

            if value is None:
                continue

            max_length = max(max_length, len(str(value)))

        hoja.column_dimensions[letra_columna].width = min(max_length + 3, 35)


def escribir_encabezados(hoja, encabezados):
    hoja.append(encabezados)

    fill = PatternFill(
        start_color="1F4E3D",
        end_color="1F4E3D",
        fill_type="solid"
    )

    for cell in hoja[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )


def aplicar_formato_numerico(hoja):
    for row in hoja.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'


def convertir_decimal(valor):
    if isinstance(valor, Decimal):
        return float(valor)

    return valor


def fecha_excel(valor):
    if not valor:
        return ""

    return valor.strftime("%d/%m/%Y")


def crear_excel(finca_activa, periodo_label, datos):
    wb = Workbook()

    hoja_resumen = wb.active
    hoja_resumen.title = "Resumen"

    resumen = datos["resumen"]

    hoja_resumen.append(["CampoSuite - Reporte consolidado"])
    hoja_resumen.append(["Finca", finca_activa.nombre])
    hoja_resumen.append(["Periodo", periodo_label])
    hoja_resumen.append([])
    hoja_resumen.append(["Indicador", "Valor"])

    resumen_items = [
        ("Total café producido kg", resumen["total_cafe_kg"]),
        ("Total café producido arrobas", resumen["total_cafe_arrobas"]),
        ("Total plátano / banano / comino kg", resumen["total_agricola_kg"]),
        ("Racimos producidos", resumen["total_racimos"]),

        ("Total ventas recibidas", resumen["total_ventas"]),
        ("Ventas café", resumen["ventas_cafe"]),
        ("Ventas plátano / banano / comino", resumen["ventas_agricolas"]),
        ("Valor esperado ventas", resumen["valor_esperado_ventas"]),
        ("Diferencia ventas", resumen["diferencia_ventas"]),

        ("Cantidad de movimientos de insumos", resumen["cantidad_movimientos"]),
        ("Valor entradas de insumos", resumen["total_entradas_valor"]),
        ("Valor salidas de insumos", resumen["total_salidas_valor"]),
        ("Ajustes registrados", resumen["total_ajustes"]),

        ("Elementos en inventario", resumen["total_elementos_inventario"]),
        ("Valor estimado inventario actual", resumen["valor_estimado_inventario"]),
    ]

    for nombre, valor in resumen_items:
        hoja_resumen.append([nombre, convertir_decimal(valor)])

    hoja_resumen["A1"].font = Font(bold=True, size=14)
    hoja_resumen["A5"].font = Font(bold=True)
    hoja_resumen["B5"].font = Font(bold=True)
    configurar_hoja(hoja_resumen)
    aplicar_formato_numerico(hoja_resumen)

    hoja_produccion = wb.create_sheet("Produccion")
    escribir_encabezados(
        hoja_produccion,
        [
            "Fecha",
            "Tipo producción",
            "Producto / cultivo",
            "Lote",
            "Kilos",
            "Arrobas",
            "Racimos",
            "Kg promedio por racimo",
            "Observaciones"
        ]
    )

    for fila in datos["filas_produccion"]:
        hoja_produccion.append([
            fecha_excel(fila["fecha"]),
            fila["tipo_produccion"],
            fila["producto"],
            fila["lote"],
            convertir_decimal(fila["kilos"]),
            convertir_decimal(fila["arrobas"]),
            fila["racimos"],
            convertir_decimal(fila["kg_promedio_racimo"])
            if fila["kg_promedio_racimo"] != ""
            else "",
            fila["observaciones"]
        ])

    configurar_hoja(hoja_produccion)
    aplicar_formato_numerico(hoja_produccion)

    hoja_ventas = wb.create_sheet("Ventas")
    escribir_encabezados(
        hoja_ventas,
        [
            "Fecha",
            "Producto",
            "Tipo",
            "Kilos vendidos",
            "Arrobas",
            "Racimos",
            "Kg promedio por racimo",
            "Precio base",
            "Valor esperado",
            "Valor recibido",
            "Diferencia",
            "Observaciones"
        ]
    )

    for fila in datos["filas_ventas"]:
        hoja_ventas.append([
            fecha_excel(fila["fecha"]),
            fila["producto"],
            fila["tipo"],
            convertir_decimal(fila["kilos"]),
            convertir_decimal(fila["arrobas"]),
            fila["racimos"],
            convertir_decimal(fila["kg_promedio_racimo"])
            if fila["kg_promedio_racimo"] != ""
            else "",
            convertir_decimal(fila["precio_unitario"]),
            convertir_decimal(fila["valor_esperado"]),
            convertir_decimal(fila["valor_recibido"]),
            convertir_decimal(fila["diferencia"]),
            fila["observaciones"]
        ])

    configurar_hoja(hoja_ventas)
    aplicar_formato_numerico(hoja_ventas)

    hoja_movimientos = wb.create_sheet("Movimientos Insumos")
    escribir_encabezados(
        hoja_movimientos,
        [
            "Fecha",
            "Elemento",
            "Categoría",
            "Tipo movimiento",
            "Cantidad",
            "Unidad",
            "Costo unitario",
            "Valor movimiento",
            "Proveedor",
            "Motivo",
            "Observaciones"
        ]
    )

    for fila in datos["filas_movimientos"]:
        hoja_movimientos.append([
            fecha_excel(fila["fecha"]),
            fila["elemento"],
            fila["categoria"],
            fila["tipo_movimiento"],
            convertir_decimal(fila["cantidad"]),
            fila["unidad"],
            convertir_decimal(fila["costo_unitario"]),
            convertir_decimal(fila["valor_movimiento"]),
            fila["proveedor"],
            fila["motivo"],
            fila["observaciones"]
        ])

    configurar_hoja(hoja_movimientos)
    aplicar_formato_numerico(hoja_movimientos)

    hoja_existencias = wb.create_sheet("Existencias Actuales")
    escribir_encabezados(
        hoja_existencias,
        [
            "Elemento",
            "Categoría",
            "Cantidad disponible",
            "Unidad",
            "Costo unitario",
            "Valor estimado",
            "Proveedor",
            "Ubicación",
            "Observaciones"
        ]
    )

    for fila in datos["filas_existencias"]:
        hoja_existencias.append([
            fila["elemento"],
            fila["categoria"],
            convertir_decimal(fila["cantidad_disponible"]),
            fila["unidad"],
            convertir_decimal(fila["costo_unitario"]),
            convertir_decimal(fila["valor_estimado"]),
            fila["proveedor"],
            fila["ubicacion"],
            fila["observaciones"]
        ])

    configurar_hoja(hoja_existencias)
    aplicar_formato_numerico(hoja_existencias)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output


@reportes_bp.route("/exportar-excel")
@login_required
def exportar_excel():
    finca_activa = obtener_finca_activa()

    anio, mes, fecha_inicio, fecha_fin, periodo_label = obtener_periodo()

    datos = obtener_datos_reporte(
        finca_activa.id,
        fecha_inicio,
        fecha_fin
    )

    archivo_excel = crear_excel(
        finca_activa,
        periodo_label,
        datos
    )

    codigo_finca = getattr(finca_activa, "codigo", "finca")

    if mes:
        nombre_archivo = f"reporte_camposuite_{codigo_finca}_{anio}_{mes:02d}.xlsx"
    else:
        nombre_archivo = f"reporte_camposuite_{codigo_finca}_{anio}.xlsx"

    return send_file(
        archivo_excel,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )